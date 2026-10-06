"""Draw the study samples (fixed seed, reproducible).

  analysis_sample.csv  2,000 reviews, 250 per app, >= 5 words   -> RQ2, RQ3 (LLM sentiment), RQ1 predictions
  gold_set.xlsx        400 of those, 50 per app, blank label columns for MANUAL annotation (RQ1)

Usage: python -m analysis.sampling
"""
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from config import ASPECTS, DATA_PROCESSED, DATA_RAW

SEED = 42
PER_APP_ANALYSIS = 250  # 8 x 250 = 2,000 >= max minimum sample size (see analysis/sample_size.py)
PER_APP_GOLD = 50       # 8 x 50 = 400 >= 385 (95% CI +/-5% on accuracy)
MIN_WORDS = 5           # shorter reviews ("good app") carry no aspect information

GOLD_XLSX = DATA_PROCESSED / "gold_set.xlsx"


def main():
    df = pd.read_csv(sorted(DATA_RAW.glob("playstore_reviews_*.csv"))[-1])
    df["n_words"] = df["text"].str.split().str.len()
    eligible = df[df["n_words"] >= MIN_WORDS]
    print("Eligible reviews per app:\n", eligible["app_name"].value_counts().to_string())

    sample = eligible.groupby("app_name").sample(n=PER_APP_ANALYSIS, random_state=SEED)
    sample.to_csv(DATA_PROCESSED / "analysis_sample.csv", index=False)

    gold = sample.groupby("app_name").sample(n=PER_APP_GOLD, random_state=SEED).sample(frac=1, random_state=SEED)
    sheet = pd.DataFrame({
        "review_id": gold["review_id"],
        "app_name": gold["app_name"],
        "text": gold["text"],
        "label_sentiment": "",       # positive / neutral / negative
        "label_relevant": "",        # yes / no
        "label_primary_aspect": "",  # one of ASPECTS, or none
        "label_language": "",        # english / hinglish / hindi / other_indian_language / other
        "notes": "",
    })
    # Star rating is deliberately NOT shown: the annotator must judge the text alone.
    sheet.to_excel(GOLD_XLSX, index=False, sheet_name="label_me")
    _format_sheet()
    print(f"\nanalysis_sample.csv: {len(sample)} rows | gold_set.xlsx: {len(sheet)} rows -> {GOLD_XLSX}")


def _format_sheet():
    wb = load_workbook(GOLD_XLSX)
    ws = wb["label_me"]
    widths = {"A": 14, "B": 16, "C": 90, "D": 16, "E": 14, "F": 30, "G": 22, "H": 30}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDE6F7")
    for row in ws.iter_rows(min_row=2):
        row[2].alignment = Alignment(wrap_text=True, vertical="top")
    n = ws.max_row
    for col, options in {
        "D": ["positive", "neutral", "negative"],
        "E": ["yes", "no"],
        "F": ASPECTS + ["none"],
        "G": ["english", "hinglish", "hindi", "other_indian_language", "other"],
    }.items():
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{n}")
    ws.freeze_panes = "D2"
    wb.save(GOLD_XLSX)


if __name__ == "__main__":
    main()
