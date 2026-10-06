"""Write the disclosed AI reference labels into gold_set.xlsx, with columns for the researcher's check.

The 400 gold-set reviews were labelled by Claude (Anthropic, model claude-opus-5-5) on 2026-09-27, following
docs/labelling_guidelines.md, blind to star ratings and to every system's predictions. The researcher then
reviews each row: `check` = ok, or `fixed` with corrected values in the corrected_* columns.
Final RQ1 ground truth = corrected value where given, otherwise the AI label.

Usage: python -m analysis.gold_ai_labels   (reads data/processed/_ai_codes.txt, one "idx s r a l" line per review)
"""
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from config import ASPECTS, DATA_PROCESSED

GOLD = DATA_PROCESSED / "gold_set.xlsx"
PROVENANCE = DATA_PROCESSED / "gold_ai_labels.csv"
ANNOTATOR = "Claude (Anthropic) claude-opus-5-5, 2026-09-27, docs/labelling_guidelines.md, blind to ratings/predictions"

SENT = {"n": "negative", "u": "neutral", "p": "positive"}
REL = {"y": "yes", "n": "no"}
ASP = {"L": "login_otp_authentication", "S": "server_downtime_performance", "U": "ui_usability",
       "K": "aadhaar_kyc_linking", "P": "payments_transactions", "D": "documents_certificates",
       "C": "customer_support_grievance", "V": "privacy_security", "B": "app_update_bugs", "O": "other", "X": "none"}
LANG = {"e": "english", "h": "hinglish", "d": "hindi", "i": "other_indian_language", "o": "other"}
LANGS = list(LANG.values())


def main():
    gold = pd.read_excel(GOLD)
    codes = {}
    for line in (DATA_PROCESSED / "_ai_codes.txt").read_text(encoding="utf-8").split("\n"):
        if line.strip():
            i, s, r, a, l = line.split()
            codes[int(i)] = (SENT[s], REL[r], ASP[a], LANG[l])
    assert sorted(codes) == list(range(len(gold))), "need exactly one label per gold row"

    out = gold[["review_id", "app_name", "text"]].copy()
    out["ai_sentiment"], out["ai_relevant"], out["ai_primary_aspect"], out["ai_language"] = zip(
        *[codes[i] for i in range(len(gold))])
    for col in ["check", "corrected_sentiment", "corrected_relevant", "corrected_primary_aspect",
                "corrected_language", "notes"]:
        out[col] = ""
    out.to_excel(GOLD, index=False, sheet_name="label_me")
    out.drop(columns=["text", "check", "corrected_sentiment", "corrected_relevant", "corrected_primary_aspect",
                      "corrected_language", "notes"]).assign(annotator=ANNOTATOR).to_csv(PROVENANCE, index=False)
    _format(len(out))
    print(out["ai_sentiment"].value_counts().to_string(), f"\n-> {GOLD}\n-> {PROVENANCE}")


def _format(n_rows: int):
    wb = load_workbook(GOLD)
    ws = wb["label_me"]
    widths = dict(zip("ABCDEFGHIJKLM", [14, 16, 80, 13, 11, 27, 18, 9, 16, 14, 27, 18, 28]))
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ai_fill, check_fill = PatternFill("solid", fgColor="E8EEF8"), PatternFill("solid", fgColor="FFF4D6")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = check_fill if cell.column >= 8 else ai_fill
    for row in ws.iter_rows(min_row=2):
        row[2].alignment = Alignment(wrap_text=True, vertical="top")
    for col, options in {"H": ["ok", "fixed"], "I": ["positive", "neutral", "negative"], "J": ["yes", "no"],
                         "K": ASPECTS + ["none"], "L": LANGS}.items():
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{n_rows + 1}")
    ws.freeze_panes = "D2"
    info = wb.create_sheet("README")
    for r, text in enumerate([
        "AI reference labels (columns D-G, blue) were produced by " + ANNOTATOR + ".",
        "Your job: read each review and its AI labels. Put 'ok' in column H if you agree.",
        "If you disagree, put 'fixed' in H and write the correct value(s) in columns I-L (only the ones that change).",
        "Final ground truth = your corrected value where given, otherwise the AI label.",
        "Definitions: docs/labelling_guidelines.md. Star ratings are deliberately not shown.",
    ], start=1):
        info.cell(row=r, column=1, value=text)
    info.column_dimensions["A"].width = 130
    wb.save(GOLD)


if __name__ == "__main__":
    main()
